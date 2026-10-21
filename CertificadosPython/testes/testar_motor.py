"""
Testes do motor com arquivos de mentira (testes/amostras.py).

Rodar:  python testes/testar_motor.py            (rápido, sem PowerPoint)
        python testes/testar_motor.py --pdf      (também gera PDF de verdade)

Os valores esperados foram calculados à mão seguindo as regras do AHK v27.
"""

import json
import os
import re
import shutil
import sys
import tempfile
import traceback

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from certificados import caminhos, coordenadores, gerador, historico, horas, planilhas, textos  # noqa: E402
from testes import amostras  # noqa: E402

TESTES = []


def teste(fn):
    TESTES.append(fn)
    return fn


def igual(obtido, esperado, contexto=""):
    assert obtido == esperado, f"{contexto}\n  esperado: {esperado!r}\n  obtido:   {obtido!r}"


# ---------------------------------------------------------------------------
# Regras de texto e horas
# ---------------------------------------------------------------------------

@teste
def regras_de_texto():
    igual(textos.chave_nome("Dra. Márcia  Souza"), "marcia souza")
    igual(textos.chave_nome("Prof. Dr. João D'Ávila"), "joao d avila")
    igual(textos.limpar_nome_docente("Dra. Maria Exemplo (cardio)"), ("Maria Exemplo", "Dra. Maria Exemplo (cardio)"))
    igual(textos.limpar_nome_docente("Maria Exemplo M"), ("Maria Exemplo", "Maria Exemplo"))
    igual(textos.sanitizar_nome_arquivo('Pós: "A/B"?.'), "Pós AB")
    igual(textos.mes_para_pasta("Abril26"), "04.26")
    igual(textos.mes_para_pasta("Março/2026"), "03.26")
    igual(textos.mes_para_pasta("Planilha1"), "Planilha1")
    igual(textos.semestre_do_mes("04.26"), "2026-1")
    igual(textos.semestre_do_mes("10.26"), "2026-2")
    igual(textos.semestre_do_mes("02.27"), "2026-2")
    igual(textos.sigla_unidade("2026-A-PTA-1S"), "PTA")
    igual(textos.separar_coordenadores("A, B / C\\D"), ["A", "B", "C", "D"])
    igual(textos.juntar_nomes(["A", "B", "C"]), "A, B ou C")
    igual(textos.email_do_docente(["clarice.lima@x.com", "paulo@x.com"], "Clarice Lima"), "clarice.lima@x.com")
    igual(textos.extrair_email("Ricardo <ricardo.alves@x.com>"), "ricardo.alves@x.com")
    assert "2026-1s" in textos.variacoes_semestre("2026-5")


@teste
def regras_de_horas():
    igual(horas.formatar_horas("08:00 - 12:00"), "4 horas")
    igual(horas.formatar_horas("13:00 às 15:30"), "2 horas e 30 minutos")
    igual(horas.formatar_horas("19h às 22h"), "3 horas")
    igual(horas.formatar_horas("4:00"), "4 horas")
    igual(horas.formatar_horas("2h30"), "2 horas e 30 minutos")
    igual(horas.formatar_horas("1 hora"), "1 hora")
    igual(horas.formatar_horas("45min"), "45 minutos")
    assert horas.formatar_horas("manhã").startswith("Horario invalido")
    igual(horas.horas_de_valor_excel(0.125), "3 horas")
    igual(horas.horas_de_valor_excel(0.104166667), "2h30")
    igual(horas.formatar_duracao_digitada("2,30"), "2 horas e 30 minutos")
    igual(horas.formatar_duracao_digitada("1"), "1 hora")
    igual(horas.corrigir_data_digitada("5/3/26"), "05/03/2026")
    igual(horas.serial_para_data(46118), "06/04/2026")


# ---------------------------------------------------------------------------
# Planilhas, coordenadores e caminhos (ambiente de mentira)
# ---------------------------------------------------------------------------

def _ambiente():
    raiz = tempfile.mkdtemp(prefix="teste_cert_")
    return raiz, amostras.criar_ambiente(raiz)


@teste
def leitura_da_mestre():
    raiz, locais = _ambiente()
    try:
        cursos = planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "2026-5")
        igual([c.posgrad for c in cursos], ["Cardiologia Clínica", "Psiquiatria Adulto", "Curso Sem Planilha"])
        igual(cursos[0].sigla_unidade, "PTA")
        igual(cursos[0].unidade_exibida, "PTA (2026-A-PTA-1S)")
        igual(planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "não existe"), None)
        igual(planilhas.listar_abas(locais.planilha_mestre), ["2026-5", "2026-6"])
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def leitura_da_planilha_do_curso():
    raiz, locais = _ambiente()
    try:
        arq = caminhos.localizar_planilha_do_curso(raiz, "Cardiologia Clínica", "2026-A-PTA-1S", "2026-5")
        assert arq.endswith("Planejamento Cardiologia Clinica PTA 2026-1S.xlsx"), arq
        igual(caminhos.localizar_planilha_do_curso(raiz, "Curso Sem Planilha", "2026-A-RJ-1S", "2026-5"), "")

        dados = planilhas.ler_dados_da_planilha(arq, ["abril26", "Maio26"])
        igual(dados.nome_pos, "Pós-Graduação em Cardiologia Clínica")
        abril, maio = dados.abas
        igual(abril.colunas, ("E", "D", "F"), "colunas detectadas na linha 3")
        resumo = [(r.linha, r.nome, r.data_aula, r.horas_formatadas, r.email) for r in abril.registros]
        igual(resumo, [
            (5, "Maria Exemplo", "06/04/2026", "4 horas", "maria.exemplo@exemplo.com"),
            (6, "João Teste", "06/04/2026", "2 horas e 30 minutos", "joao@exemplo.com"),
            # Mesma quantidade de nomes e e-mails: pareados NA ORDEM (regra do AHK).
            (8, "Paulo Coelho", "07/04/2026", "4 horas", "clarice.lima@exemplo.com"),
            (8, "Clarice Lima", "07/04/2026", "4 horas", "paulo.coelho@exemplo.com"),
            # Data exibida como "8-abr": vem do número do Excel (o AHK daria 09/04).
            (9, "Maria Exemplo", "08/04/2026", "3 horas", "sem email"),
            (10, "Pedro Sem Aula", "09/04/2026", "2 horas e 30 minutos", "pedro@exemplo.com"),
        ])
        r5 = abril.registros[0]
        assert r5.mesclado and r5.suspeita_nome_com_extra and not r5.tem_erro
        assert abril.registros[-1].erro_curso
        igual([(r.nome, r.horas_formatadas, r.email) for r in maio.registros],
              [("Ricardo Alves", "3 horas", "ricardo.alves@exemplo.com")])
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def coordenadores_e_cadastro():
    raiz, locais = _ambiente()
    try:
        cad = coordenadores.Cadastro(locais.assinaturas, log=lambda m: None)
        ana = coordenadores.resolver("Dra. Ana Beatriz Souza", cad, log=lambda m: None)
        igual((ana.nome, ana.titulo, ana.origem, ana.pendente), ("Dra. Ana Beatriz Souza", "Dra.", "pelo nome", ""))
        igual(ana.assina_como, "Dra. Ana Beatriz Souza")       # sem "Dra. Dra."
        igual(textos.sem_titulo("Prof. Dr. João Silva"), "João Silva")
        igual(ana.cargo, "Coordenadora do Curso")
        psi = coordenadores.resolver("Carlos Mendes / Paula Ribeiro", cad, log=lambda m: None)
        igual((psi.nome, psi.pendente), ("Paula Ribeiro", "titulo"))
        coordenadores.definir_titulo([ana, psi], cad, "Paula Ribeiro", "dra")
        igual((psi.titulo, psi.pendente), ("Dra.", ""))
        sem = coordenadores.resolver("Fulano Sem Assinatura", cad, log=lambda m: None)
        igual(sem.pendente, "assinatura")
        img = os.path.join(locais.assinaturas, "Ana Beatriz Souza.png")
        coordenadores.definir_assinatura(sem, cad, img)
        igual((sem.nome, sem.pendente), ("Fulano Sem Assinatura", "titulo"))
        # O cadastro foi gravado no formato do AHK e é lido de volta.
        cad2 = coordenadores.Cadastro(locais.assinaturas, log=lambda m: None)
        igual(cad2.obter("paula ribeiro")["titulo"], "Dra.")
        igual(cad2.obter("Fulano Sem Assinatura")["assinatura"], "Ana Beatriz Souza.png")
        with open(cad.arquivo, encoding="utf-8-sig") as f:
            assert "Paula Ribeiro | Dra. | " in f.read()
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


# ---------------------------------------------------------------------------
# Geração completa
# ---------------------------------------------------------------------------

class ConversorFalso:
    """Não gera PDF de verdade: só copia o .pptx com nome .pdf."""
    nome = "falso"

    def converter(self, pptx, pdf):
        shutil.copy2(pptx, pdf)


def _lote(locais, raiz):
    cad = coordenadores.Cadastro(locais.assinaturas, log=lambda m: None)
    curso = planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "2026-5")[0]
    c = gerador.CursoNoLote(curso)
    c.coordenacao = coordenadores.resolver(curso.coordenacao, cad, log=lambda m: None)
    c.planilha = caminhos.localizar_planilha_do_curso(raiz, curso.posgrad, curso.unidade, "2026-5")
    c.abas = ["Abril26"]
    c.dados = planilhas.ler_dados_da_planilha(c.planilha, c.abas)
    return c


def _gerar(modo_nome, conversor, raiz, locais):
    modo = gerador.MODOS[modo_nome]
    modo = gerador.Modo(**{**modo.__dict__, "pasta_local": os.path.join(raiz, "saida_" + modo_nome)})
    c = _lote(locais, raiz)
    problemas, _, _ = gerador.validar([c])
    for p in problemas:          # igual ao "Pular" da tela de problemas
        p.registro.pular = p.registro.tem_erro
    logs = []
    log = gerador.Log(modo_nome, pasta=os.path.join(raiz, "logs"))
    rel = gerador.gerar([c], locais, locais.modelo, "20/10/2026", modo, conversor, log,
                        email_teste="eu@teste.com")
    return c, rel, modo


@teste
def geracao_oficial():
    raiz, locais = _ambiente()
    try:
        c, rel, modo = _gerar("oficial", ConversorFalso(), raiz, locais)
        assert not rel.falhas, rel.falhas
        igual(rel.total_gerados, 5)
        envio = os.path.join(locais.declaracoes, "Para Enviar", "Cardiologia Clínica (2026-A-PTA-1S)", "04.26")
        nomes = sorted(os.listdir(envio))
        igual(nomes, ["Clarice Lima (07.04.2026).pdf", "João Teste (06.04.2026).pdf",
                      "Maria Exemplo (06.04.2026).pdf", "Maria Exemplo (08.04.2026).pdf",
                      "Paulo Coelho (07.04.2026).pdf", "manifesto.txt"])
        with open(os.path.join(envio, "manifesto.txt"), encoding="utf-8") as f:
            m = json.load(f)
        igual(list(m)[:4], ["versaoManifesto", "loteId", "dataGeracao", "curso"])
        igual((m["curso"], m["local"], m["nomeAba"], m["semestre"], m["pastaMes"]),
              ("Cardiologia Clínica", "PTA", "Abril26", "2026-1", "04.26"))
        igual(m["pastaDestino"], "2026-1/PTA/Cardiologia Clínica/04.26")
        igual(len(m["arquivosDoLote"]), 5)
        maria = next(d for d in m["certificados"] if d["docente"] == "Maria Exemplo")
        igual([a["nomeArquivo"] for a in maria["arquivos"]],
              ["Maria Exemplo (06.04.2026).pdf", "Maria Exemplo (08.04.2026).pdf"])
        igual(maria["email"], "maria.exemplo@exemplo.com")
        docentes_local = os.path.join(modo.pasta_local, "Cardiologia Clínica (2026-A-PTA-1S)", "Docentes")
        assert os.path.isdir(os.path.join(docentes_local, "Maria Exemplo"))
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def geracao_teste_e_simulacao():
    raiz, locais = _ambiente()
    try:
        _, rel, _ = _gerar("teste", ConversorFalso(), raiz, locais)
        envio = os.path.join(locais.declaracoes, "Para Enviar", "TESTE - Cardiologia Clínica (2026-A-PTA-1S)", "04.26")
        assert "TESTE - Maria Exemplo (06.04.2026).pdf" in os.listdir(envio)
        with open(os.path.join(envio, "manifesto.txt"), encoding="utf-8") as f:
            m = json.load(f)
        igual({d["email"] for d in m["certificados"]}, {"eu@teste.com"})
        assert m["pastaDestino"].startswith("TESTE/2026-1/PTA/")

        antes = set(os.listdir(os.path.join(locais.declaracoes, "Para Enviar")))
        _, rel, modo = _gerar("simulacao", ConversorFalso(), raiz, locais)
        igual(set(os.listdir(os.path.join(locais.declaracoes, "Para Enviar"))), antes, "simulação não pode tocar na nuvem")
        pasta = os.path.join(modo.pasta_local, "Cardiologia Clínica (2026-A-PTA-1S)", "04.26")
        assert "manifesto.txt" in os.listdir(pasta)
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def certificado_preenchido():
    """O .pptx sai sem nenhum {MARCADOR}, com logo e assinatura trocados."""
    from pptx import Presentation
    from certificados import modelo
    raiz, locais = _ambiente()
    try:
        base = modelo.criar_modelo_base(locais.modelo, "Dra. Ana Beatriz Souza", "Coordenadora do Curso",
                                        "20/10/2026", os.path.join(locais.assinaturas, "Ana Beatriz Souza.png"),
                                        locais.logo)
        destino = os.path.join(raiz, "c.pptx")
        modelo.preencher_certificado(base, "Maria Exemplo", "Pós em Teste", "Anatomia", "06/04/2026", "4 horas", destino)
        pres = Presentation(destino)
        texto = "\n".join(f.text_frame.text for f in pres.slides[0].shapes if f.has_text_frame)
        assert "{" not in texto, texto
        assert "Declaramos que Maria Exemplo ministrou a aula Anatomia da Pós em Teste, em 06/04/2026, com carga horária de 4 horas." in texto
        assert "São Paulo, 20/10/2026" in texto
        nomes = [f.name for f in pres.slides[0].shapes]
        assert "ASSINATURA_REAL" in nomes and "LOGO_EINSTEIN" in nomes, nomes
        coord = next(f for f in pres.slides[0].shapes if "Dra. Ana" in (f.text_frame.text if f.has_text_frame else ""))
        assinatura = next(f for f in pres.slides[0].shapes if f.name == "ASSINATURA_REAL")
        igual(assinatura.left, coord.left)
        run = coord.text_frame.paragraphs[0].runs[0]
        igual((run.text, run.font.size.pt, run.font.bold), ("Dra. Ana Beatriz Souza", 10.0, False))
        run = coord.text_frame.paragraphs[1].runs[0]
        igual((run.text, run.font.size.pt, run.font.bold), ("Coordenadora do Curso", 10.0, True))
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def historico_de_envios():
    raiz, locais = _ambiente()
    try:
        cursos = planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "2026-5")
        cardio = cursos[0]
        igual(historico.situacao(locais, cursos)[cardio.linha], [])
        c, rel, _ = _gerar("oficial", ConversorFalso(), raiz, locais)
        assert historico.registrar(locais, rel, log=lambda m: None)
        # Gerado e ainda na Para Enviar: amarelo.
        meses = historico.situacao(locais, cursos)[cardio.linha]
        igual([(m.rotulo, m.pendente) for m in meses], [("Abr26", True)])
        # O fluxo arquivou e apagou a pasta da Para Enviar: verde, com a data do envio.
        destino = os.path.join(locais.declaracoes, "2026-1", "PTA", "Cardiologia Clínica", "04.26")
        shutil.move(os.path.join(locais.declaracoes, "Para Enviar", c.nome_com_unidade, "04.26"), destino)
        meses = historico.situacao(locais, cursos)[cardio.linha]
        igual([(m.rotulo, m.pendente) for m in meses], [("Abr26", False)])
        assert re.match(r"\d\d/\d\d$", meses[0].data), meses[0].data
        # Pasta arquivada sem histórico (envio do AHK) também conta; outro curso não é afetado.
        os.makedirs(os.path.join(locais.declaracoes, "2026-1", "PTA", "Cardiologia Clínica", "03.26"))
        with open(os.path.join(locais.declaracoes, "2026-1", "PTA", "Cardiologia Clínica", "03.26", "a.pdf"), "w"):
            pass
        sit = historico.situacao(locais, cursos)
        igual([m.rotulo for m in sit[cardio.linha]], ["Mar26", "Abr26"])
        igual(sit[cursos[1].linha], [])
        igual(historico.meses_conhecidos(sit), ["04.26", "03.26"])
        igual(historico.rotulo_mes("10.26"), "Out26")
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


def geracao_com_pdf_de_verdade(pasta_saida):
    """Gera os PDFs com o PowerPoint e confere o texto de dentro deles."""
    from pypdf import PdfReader
    from certificados.pdf import ConversorPowerPoint
    raiz, locais = _ambiente()
    try:
        with ConversorPowerPoint() as conv:
            _, rel, modo = _gerar("simulacao", conv, raiz, locais)
        pasta = os.path.join(modo.pasta_local, "Cardiologia Clínica (2026-A-PTA-1S)", "04.26")
        pdf = os.path.join(pasta, "Maria Exemplo (06.04.2026).pdf")
        texto = " ".join(PdfReader(pdf).pages[0].extract_text().split())
        assert "Maria Exemplo" in texto and "Anatomia do Coração" in texto and "4 horas" in texto, texto
        os.makedirs(pasta_saida, exist_ok=True)
        for n in os.listdir(pasta):
            shutil.copy2(os.path.join(pasta, n), pasta_saida)
        print("   PDFs copiados para", pasta_saida)
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


if __name__ == "__main__":
    falhas = 0
    for fn in TESTES:
        try:
            fn()
            print(f"OK     {fn.__name__}")
        except Exception:
            falhas += 1
            print(f"FALHOU {fn.__name__}\n{traceback.format_exc()}")
    if "--pdf" in sys.argv:
        try:
            geracao_com_pdf_de_verdade(sys.argv[sys.argv.index("--pdf") + 1] if len(sys.argv) > sys.argv.index("--pdf") + 1
                                       else os.path.join(tempfile.gettempdir(), "certificados_pdf_teste"))
            print("OK     geracao_com_pdf_de_verdade")
        except Exception:
            falhas += 1
            print(f"FALHOU geracao_com_pdf_de_verdade\n{traceback.format_exc()}")
    print(f"\n{len(TESTES) + ('--pdf' in sys.argv) - falhas} de {len(TESTES) + ('--pdf' in sys.argv)} passaram")
    sys.exit(1 if falhas else 0)
