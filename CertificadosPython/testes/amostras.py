"""
Arquivos de mentira com a mesma estrutura dos reais, para testar o motor sem
usar dados de docentes: OneDrive falso com "Planilha compartilhada -
Declarações", planilha mestre, planilha de um curso, modelo .pptx, logo e
assinaturas.
"""

import os
import struct
import zlib
from datetime import date, datetime, time

import openpyxl
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

from certificados import caminhos


def png(largura, altura, rgb):
    linha = b"\x00" + bytes(rgb) * largura
    def bloco(tipo, dados):
        return struct.pack(">I", len(dados)) + tipo + dados + struct.pack(">I", zlib.crc32(tipo + dados) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + bloco(b"IHDR", struct.pack(">IIBBBBB", largura, altura, 8, 2, 0, 0, 0))
            + bloco(b"IDAT", zlib.compress(linha * altura)) + bloco(b"IEND", b""))


def _salvar_png(caminho, larg, alt, rgb):
    with open(caminho, "wb") as f:
        f.write(png(larg, alt, rgb))


def criar_modelo(caminho, pasta_imagens):
    """Modelo parecido com o real: logo no topo, texto com {MARCADORES}
    quebrados em pedaços, assinatura acima do coordenador."""
    pres = Presentation()
    pres.slide_width, pres.slide_height = Cm(29.7), Cm(21)
    slide = pres.slides.add_slide(pres.slide_layouts[6])
    logo_modelo = os.path.join(pasta_imagens, "logo_modelo.png")
    assin_modelo = os.path.join(pasta_imagens, "assinatura_modelo.png")
    _salvar_png(logo_modelo, 200, 120, (180, 180, 180))
    _salvar_png(assin_modelo, 200, 80, (220, 220, 220))
    slide.shapes.add_picture(logo_modelo, Cm(11), Cm(0.8), Cm(7), Cm(4))

    def caixa(pedacos, topo, tamanho, negrito=False, altura=Cm(1.5)):
        tb = slide.shapes.add_textbox(Cm(2), topo, Cm(25.7), altura)
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        for texto in pedacos:
            r = p.add_run()
            r.text = texto
            r.font.size = Pt(tamanho)
            r.font.bold = negrito
            r.font.name = "Calibri"
        return tb

    caixa(["DECLARAÇÃO"], Cm(5.2), 30, True)
    caixa(["Declaramos que ", "{NO", "ME}", " ministrou a aula ", "{CURSO}", " da ", "{POSGRADUACAO}",
           ", em ", "{DATA}", ", com carga horária de ", "{HO", "RAS}", "."], Cm(7.5), 16, altura=Cm(3))
    caixa(["São Paulo, {DATAEMISSAO}"], Cm(12), 12)
    slide.shapes.add_picture(assin_modelo, Cm(12), Cm(14), Cm(5), Cm(2))
    tb = slide.shapes.add_textbox(Cm(10), Cm(16.5), Cm(9.7), Cm(1.6))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "{COORDENADOR}"
    r.font.size = Pt(14)
    p2 = tb.text_frame.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "{CARGO}"
    r2.font.size = Pt(12)
    pres.save(caminho)


def criar_mestre(caminho):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026-5"
    ws.append(["Coordenação", "Pós-graduação", "Código da turma", "Turma"])
    ws.append(["Dra. Ana Beatriz Souza", "Cardiologia Clínica", "2026-A-PTA-1S", "T1"])
    ws.append(["Carlos Mendes / Paula Ribeiro", "Psiquiatria Adulto", "2026-A-CEP-1S", "T2"])
    ws.append(["Fulano Sem Assinatura", "Curso Sem Planilha", "2026-A-RJ-1S", "T3"])
    wb.create_sheet("2026-6").append(["Coordenação", "Pós-graduação"])
    wb.save(caminho)


def criar_planilha_curso(caminho, nome_pos="Pós-Graduação em Cardiologia Clínica"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Abril26"
    ws["A2"] = nome_pos
    for col, titulo in enumerate(["Data", "Horário", "Sala", "Disciplina / Aula", "Docente", "E-mail"], 1):
        ws.cell(row=3, column=col, value=titulo)
    linhas = [
        (date(2026, 4, 6), "08:00 - 12:00", "M1", "Anatomia do Coração", "Dra. Maria Exemplo (cardio)", "maria.exemplo@exemplo.com"),
        (None, "13:00 às 15:30", "M1", "Fisiologia", "Dr João Teste", "joao@exemplo.com"),
        (None, "15:30 - 16:00", "", "Coffee break", "", ""),
        (date(2026, 4, 7), time(4, 0), "M2", "Eletrocardiograma", "Paulo Coelho / Clarice Lima",
         "clarice.lima@exemplo.com; paulo.coelho@exemplo.com"),
        (date(2026, 4, 8), 0.125, "M2", "Ecocardiograma", "Maria Exemplo M", "sem email"),
        (datetime(2026, 4, 9), "2h30", "M3", "", "Pedro Sem Aula", "pedro@exemplo.com"),
    ]
    for i, linha in enumerate(linhas, start=5):
        for col, valor in enumerate(linha, 1):
            if valor is not None:
                ws.cell(row=i, column=col, value=valor)
    ws.cell(row=5, column=1).number_format = "dd/mm/yyyy"
    ws.cell(row=8, column=1).number_format = "dd/mm/yyyy"
    ws.cell(row=8, column=2).number_format = "h:mm"
    ws.cell(row=9, column=1).number_format = "d-mmm"           # texto "8-abr": cai no caminho do número
    ws.cell(row=10, column=1).number_format = "dd/mm/yyyy"
    ws.merge_cells("A5:A7")                                      # data mesclada em 3 linhas
    ws.cell(row=12, column=1, value="Coordenação: Dra. Ana Beatriz Souza")   # rodapé

    ws2 = wb.create_sheet("Maio26")
    ws2["A2"] = nome_pos
    for col, titulo in enumerate(["Data", "Horário", "Sala", "Aula", "Professor", "Contato"], 1):
        ws2.cell(row=3, column=col, value=titulo)
    ws2.append([])
    ws2.cell(row=5, column=1, value="04/05/2026")
    ws2.cell(row=5, column=2, value="19h às 22h")
    ws2.cell(row=5, column=4, value="Insuficiência Cardíaca")
    ws2.cell(row=5, column=5, value="Prof. Ricardo Alves")
    ws2.cell(row=5, column=6, value="Ricardo <ricardo.alves@exemplo.com>")
    wb.save(caminho)


def criar_ambiente(raiz):
    """Cria a árvore inteira dentro de "raiz" (que faz o papel do OneDrive)."""
    decl = os.path.join(raiz, caminhos.NOME_DECLARACOES)
    assin = os.path.join(decl, caminhos.NOME_ASSINATURAS)
    pos = os.path.join(raiz, "Planilha compartilhada - Pós-graduação 2026-1S", "Cardiologia")
    for p in (decl, assin, pos, os.path.join(decl, "Para Enviar")):
        os.makedirs(p, exist_ok=True)
    criar_mestre(os.path.join(decl, caminhos.NOME_MESTRE))
    criar_modelo(os.path.join(decl, caminhos.NOME_MODELO), assin)
    criar_planilha_curso(os.path.join(pos, "Planejamento Cardiologia Clinica PTA 2026-1S.xlsx"))
    _salvar_png(os.path.join(assin, "Ana Beatriz Souza.png"), 260, 100, (30, 60, 160))
    _salvar_png(os.path.join(assin, "Dra Paula A. Ribeiro.png"), 260, 100, (160, 30, 60))
    _salvar_png(os.path.join(assin, caminhos.NOME_LOGO), 320, 200, (15, 107, 111))
    with open(os.path.join(assin, "coordenadores.txt"), "w", encoding="utf-8-sig") as f:
        f.write("# teste\r\nAna Beatriz Souza | Dra. | \r\n")
    return caminhos.Locais.descobrir(raiz)
