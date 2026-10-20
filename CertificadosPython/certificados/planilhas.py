"""
Leitura da planilha mestre e das planilhas de curso (openpyxl, sem Excel).

Mesmas regras de LerCursosDaMestre / LerDadosDaPlanilha / LerRegistro do AHK.
O AHK lê às vezes o VALOR da célula e às vezes o TEXTO que o Excel mostra
(.Text). O openpyxl só entrega o valor, então texto_exibido() refaz o texto
a partir do formato de número da célula.
"""

import copy
import os
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

from . import textos, horas

# ---------------------------------------------------------------------------
# Texto exibido pelo Excel (.Text)
# ---------------------------------------------------------------------------

_MESES_CURTOS = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
_MESES_LONGOS = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
                 "setembro", "outubro", "novembro", "dezembro"]
_DIAS_CURTOS = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]
_DIAS_LONGOS = ["segunda-feira", "terça-feira", "quarta-feira", "quinta-feira", "sexta-feira", "sábado", "domingo"]
_TOKENS = re.compile(r"\[h+\]|\[m+\]|yyyy|yy|mmmmm|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|am/pm|a/p|.", re.I)


def _limpar_formato(fmt):
    fmt = (fmt or "General").split(";")[0]
    fmt = re.sub(r"\[\$[^\]]*\]", "", fmt)          # [$-416] (idioma)
    fmt = re.sub(r"\[(?!h+\]|m+\]|s+\])[^\]]*\]", "", fmt, flags=re.I)  # [Red] etc.
    return fmt


def _formato_de_data(fmt):
    sem_literais = re.sub(r'"[^"]*"|\\.', "", fmt)
    return re.search(r"[dmyhs]", sem_literais, re.I) is not None


def _como_datetime(valor):
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime(valor.year, valor.month, valor.day)
    if isinstance(valor, time):
        return datetime(1899, 12, 30, valor.hour, valor.minute, valor.second)
    if isinstance(valor, (int, float)):
        return datetime(1899, 12, 30) + timedelta(days=float(valor))
    return None


def _formatar_data_hora(valor, fmt):
    total = horas.para_serial(valor)
    dt = _como_datetime(valor if not isinstance(valor, timedelta) else total)
    if dt is None:
        return None
    tokens = _TOKENS.findall(re.sub(r"\\(.)", r"\1", fmt.replace('"', "")))
    tem_ampm = any(t.lower() in ("am/pm", "a/p") for t in tokens)
    saida = []
    for i, t in enumerate(tokens):
        tl = t.lower()
        # "m" depois de hora ou antes de segundo = minuto.
        anterior = next((x.lower() for x in reversed(tokens[:i]) if re.fullmatch(r"\[?[hdmys]+\]?", x, re.I)), "")
        seguinte = next((x.lower() for x in tokens[i + 1:] if re.fullmatch(r"\[?[hdmys]+\]?", x, re.I)), "")
        if tl in ("m", "mm") and (anterior.strip("[]").startswith("h") or seguinte.startswith("s")):
            saida.append(f"{dt.minute:02d}" if tl == "mm" else str(dt.minute))
        elif tl.startswith("[h"):
            saida.append(str(int(total * 24)).zfill(len(tl) - 2))
        elif tl.startswith("[m"):
            saida.append(str(int(total * 1440)).zfill(len(tl) - 2))
        elif tl == "yyyy":
            saida.append(f"{dt.year:04d}")
        elif tl == "yy":
            saida.append(f"{dt.year % 100:02d}")
        elif tl == "mmmmm":
            saida.append(_MESES_LONGOS[dt.month - 1][0])
        elif tl == "mmmm":
            saida.append(_MESES_LONGOS[dt.month - 1])
        elif tl == "mmm":
            saida.append(_MESES_CURTOS[dt.month - 1])
        elif tl == "mm":
            saida.append(f"{dt.month:02d}")
        elif tl == "m":
            saida.append(str(dt.month))
        elif tl == "dddd":
            saida.append(_DIAS_LONGOS[dt.weekday()])
        elif tl == "ddd":
            saida.append(_DIAS_CURTOS[dt.weekday()])
        elif tl == "dd":
            saida.append(f"{dt.day:02d}")
        elif tl == "d":
            saida.append(str(dt.day))
        elif tl in ("hh", "h"):
            h = (dt.hour % 12 or 12) if tem_ampm else dt.hour
            saida.append(f"{h:02d}" if tl == "hh" else str(h))
        elif tl in ("ss", "s"):
            saida.append(f"{dt.second:02d}" if tl == "ss" else str(dt.second))
        elif tl == "am/pm":
            saida.append("AM" if dt.hour < 12 else "PM")
        elif tl == "a/p":
            saida.append("A" if dt.hour < 12 else "P")
        else:
            saida.append(t)
    return "".join(saida)


def _numero_geral(valor):
    """Número em formato Geral, como o Excel em português mostra (vírgula decimal)."""
    if isinstance(valor, bool):
        return "VERDADEIRO" if valor else "FALSO"
    if isinstance(valor, int) or float(valor).is_integer():
        return str(int(valor))
    return f"{valor:.10g}".replace(".", ",")


def texto_exibido(celula):
    """O que o Excel mostraria na célula (equivalente ao .Text do AHK)."""
    valor = celula.value
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor
    fmt = _limpar_formato(celula.number_format)
    if isinstance(valor, (datetime, date, time, timedelta)) or (
            isinstance(valor, (int, float)) and not isinstance(valor, bool) and _formato_de_data(fmt)):
        texto = _formatar_data_hora(valor, fmt if _formato_de_data(fmt) else "dd/mm/yyyy")
        if texto is not None:
            return texto
    if isinstance(valor, (int, float)):
        return _numero_geral(valor)
    return str(valor)


def valor_texto(valor):
    """String(celula.Value) do AHK, para células de texto (nome, aula, e-mail...)."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


# ---------------------------------------------------------------------------
# Aba com células mescladas resolvidas
# ---------------------------------------------------------------------------

class Aba:
    """Uma aba lida pelo openpyxl, com a mesma ideia do ResolverCelulaMesclada."""

    def __init__(self, ws):
        self.ws = ws
        self.nome = ws.title
        self._origem = {}      # (linha, coluna) -> (linha, coluna) da célula de cima/esquerda
        for faixa in ws.merged_cells.ranges:
            for r in range(faixa.min_row, faixa.max_row + 1):
                for c in range(faixa.min_col, faixa.max_col + 1):
                    self._origem[(r, c)] = (faixa.min_row, faixa.min_col)

    def mesclada(self, linha, coluna):
        return (linha, coluna) in self._origem

    def celula(self, linha, coluna):
        """Célula direta (sem resolver mescla)."""
        return self.ws.cell(row=linha, column=coluna)

    def celula_resolvida(self, linha, coluna):
        r, c = self._origem.get((linha, coluna), (linha, coluna))
        return self.ws.cell(row=r, column=c)

    def ultima_linha(self, coluna):
        """Cells(Rows.Count, col).End(xlUp).Row: última linha com algo na coluna."""
        for r in range(self.ws.max_row, 0, -1):
            v = self.ws.cell(row=r, column=coluna).value
            if v is not None and str(v) != "":
                return r
        return 1

    def max_colunas(self):
        return max(50, self.ws.max_column or 0)


def _abrir(caminho):
    """Abre uma CÓPIA da planilha (o original no OneDrive nunca fica preso)."""
    pasta = tempfile.mkdtemp(prefix="cert_planilha_")
    copia = os.path.join(pasta, os.path.basename(caminho))
    shutil.copy2(caminho, copia)
    try:
        return openpyxl.load_workbook(copia, data_only=True)
    finally:
        shutil.rmtree(pasta, ignore_errors=True)


def _achar_aba(wb, nome):
    """ObterAba: compara sem espaço nas pontas e sem diferenciar maiúscula."""
    for ws in wb.worksheets:
        if ws.title.strip().casefold() == str(nome).strip().casefold():
            return ws
    return None


def listar_abas(caminho):
    return [ws.title for ws in _abrir(caminho).worksheets]


# ---------------------------------------------------------------------------
# Planilha mestre
# ---------------------------------------------------------------------------

@dataclass
class Curso:
    linha: int
    coordenacao: str
    posgrad: str
    unidade: str          # código da turma, ex. "2026-A-PTA-1S"
    sigla_unidade: str    # ex. "PTA"
    turma: str

    @property
    def unidade_exibida(self):
        if self.sigla_unidade and self.sigla_unidade != self.unidade:
            return f"{self.sigla_unidade} ({self.unidade})"
        return self.unidade


def ler_cursos_da_mestre(caminho, nome_aba):
    """LerCursosDaMestre: A = coordenação, B = pós, C = código da turma, D = turma.
    Devolve None se a aba não existir."""
    ws = _achar_aba(_abrir(caminho), nome_aba)
    if ws is None:
        return None
    aba = Aba(ws)
    cursos = []
    for linha in range(2, aba.ultima_linha(2) + 1):
        valor = lambda col: valor_texto(aba.celula(linha, col).value).strip()
        posgrad = valor(2)
        if not posgrad:
            continue
        codigo = valor(3)
        cursos.append(Curso(linha, valor(1), posgrad, codigo, textos.sigla_unidade(codigo), valor(4)))
    return cursos


# ---------------------------------------------------------------------------
# Planilha do curso
# ---------------------------------------------------------------------------

@dataclass
class Registro:
    aba: str
    linha: int
    nome: str
    curso: str
    data_aula: str
    horario_txt: str
    horas_formatadas: str
    nome_original_bruto: str
    email: str = ""
    email_valido: bool = False
    email_bruto: str = ""
    mesclado: bool = False
    rodape: bool = False
    erro_nome: bool = False
    erro_curso: bool = False
    erro_data: bool = False
    erro_horas: bool = False
    suspeita_nome_email: bool = False
    suspeita_curso_data: bool = False
    suspeita_nome_com_extra: bool = False
    linha_com_varios_docentes: str = ""
    pular: bool = False

    @property
    def tem_erro(self):
        return self.erro_nome or self.erro_curso or self.erro_data or self.erro_horas

    @property
    def tem_problema(self):
        return (self.tem_erro or self.suspeita_nome_email or self.suspeita_curso_data
                or self.suspeita_nome_com_extra)

    # Correções manuais (CorrigirNomeRegistro, CorrigirCursoRegistro...)
    def corrigir_nome(self, valor):
        self.nome, self.nome_original_bruto = textos.limpar_nome_docente(valor)
        self.erro_nome = self.nome == ""
        self.suspeita_nome_email = "@" in self.nome
        self.suspeita_nome_com_extra = False

    def corrigir_curso(self, valor):
        self.curso = str(valor).strip()
        self.erro_curso = self.curso == ""
        self.suspeita_curso_data = bool(self.curso) and re.match(r"^\d{1,2}[/:h]\d", self.curso) is not None

    def corrigir_data(self, valor):
        data = horas.corrigir_data_digitada(valor)
        if data:
            self.data_aula, self.erro_data = data, False
        return bool(data)

    def corrigir_horas(self, valor):
        duracao = horas.formatar_duracao_digitada(valor)
        if duracao:
            self.horas_formatadas = self.horario_txt = duracao
            self.erro_horas = False
        return bool(duracao)


def _coluna_por_termos(aba, termos):
    for c in range(1, aba.max_colunas() + 1):
        celula = aba.celula(3, c)
        texto = valor_texto(celula.value) or texto_exibido(celula)
        norm = textos.normalizar(texto)
        if any(t in norm for t in termos):
            return c
    return None


def detectar_colunas(aba):
    """DetectarColunasPorAba: procura os títulos na linha 3. Padrão: E, D, F."""
    nome = _coluna_por_termos(aba, ["docente", "professor"]) or column_index_from_string("E")
    curso = _coluna_por_termos(aba, ["disciplina", "aula", "modulo", "conteudo", "tema", "materia", "topico"]) \
        or column_index_from_string("D")
    email = _coluna_por_termos(aba, ["email", "e-mail", "contato", "correio", "eletronico"]) \
        or column_index_from_string("F")
    return nome, curso, email


def _ler_data(celula):
    """LerDataCelula: texto exibido; senão número do Excel; senão procura no texto."""
    txt = texto_exibido(celula).strip()
    if horas.data_valida(txt):
        return txt
    serial = horas.para_serial(celula.value)
    if isinstance(serial, (int, float)) and not isinstance(serial, bool) and serial > 1:
        convertida = horas.serial_para_data(serial)
        if horas.data_valida(convertida):
            return convertida
    v = valor_texto(celula.value).strip()
    if horas.data_valida(v):
        return v
    m = re.search(r"\d{1,2}/\d{1,2}/\d{2,4}", v)
    return m.group(0) if m else ""


_IGNORAR = ["almoco", "intervalo", "coffee break", "coffeebreak", "lanche", "refeicao", "pausa", "descanso"]


def ler_registro(aba, linha, col_nome, col_curso, col_email):
    """LerRegistro."""
    cel_data = aba.celula_resolvida(linha, 1)
    cel_horario = aba.celula_resolvida(linha, 2)
    cel_nome = aba.celula_resolvida(linha, col_nome)
    cel_curso = aba.celula_resolvida(linha, col_curso)
    mesclado = any(aba.mesclada(linha, c) for c in (1, 2, col_nome, col_curso))

    data_bruta = _ler_data(cel_data)
    data_aula = data_bruta if horas.data_valida(data_bruta) else ""

    nome, bruto = textos.limpar_nome_docente(valor_texto(cel_nome.value))
    curso = valor_texto(cel_curso.value).strip()

    horario_txt = texto_exibido(cel_horario)
    if horario_txt == "" or re.match(r"^\d[.,]\d", horario_txt):
        valor = horas.para_serial(cel_horario.value)
        horas_fmt = horas.horas_de_valor_excel(valor)
    else:
        horas_fmt = horas.formatar_horas(horario_txt)

    email_bruto = valor_texto(aba.celula(linha, col_email).value).strip()   # sem resolver mescla, como no AHK
    email = textos.extrair_email(email_bruto)

    reg = Registro(
        aba=aba.nome, linha=linha, nome=nome, curso=curso, data_aula=data_aula,
        horario_txt=horario_txt, horas_formatadas=horas_fmt, nome_original_bruto=bruto,
        email=email, email_valido=textos.email_valido(email), email_bruto=email_bruto, mesclado=mesclado,
    )
    reg.erro_nome = nome == ""
    reg.erro_curso = curso == ""
    reg.erro_data = data_aula == ""
    reg.erro_horas = "invalido" in horas_fmt.lower()
    reg.suspeita_nome_email = "@" in nome
    reg.suspeita_curso_data = bool(curso) and re.match(r"^\d{1,2}[/:h]\d", curso) is not None
    reg.suspeita_nome_com_extra = nome != re.sub(r"\s+", " ", bruto).strip()

    preenchidos = sum([not reg.erro_nome, not reg.erro_curso, not reg.erro_data, horario_txt != ""])
    normalizados = [textos.normalizar(x) for x in (nome, curso, horario_txt)]
    intervalo = any(t in n for t in _IGNORAR for n in normalizados)
    reg.rodape = preenchidos <= 1 or (nome == "" and curso == "") or intervalo
    return reg


def dividir_por_docente(reg):
    """DividirRegistroPorDocente: "Dr Paulo / Dra Clarice" vira um registro por nome."""
    sem_parenteses = re.sub(r"\([^)]*\)", "", reg.nome_original_bruto)
    partes = [p.strip() for p in re.split(r"[/;&+|\n\r]", sem_parenteses) if p.strip()]
    if len(partes) < 2:
        return [reg]
    emails = textos.todos_os_emails(reg.email_bruto)
    novos = []
    for i, parte in enumerate(partes):
        novo = copy.copy(reg)
        novo.corrigir_nome(parte)
        novo.linha_com_varios_docentes = reg.nome_original_bruto
        email = emails[i] if len(emails) == len(partes) else textos.email_do_docente(emails, novo.nome)
        novo.email, novo.email_valido = email, textos.email_valido(email)
        novos.append(novo)
    return novos


@dataclass
class DadosAba:
    aba: str
    registros: list
    colunas: tuple = ()


@dataclass
class DadosCurso:
    nome_pos: str
    abas: list = field(default_factory=list)


def ler_dados_da_planilha(caminho, abas):
    """LerDadosDaPlanilha: A2 = nome da pós; linhas da 5 até a última data."""
    wb = _abrir(caminho)
    primeira = _achar_aba(wb, abas[0])
    nome_pos = valor_texto(primeira["A2"].value).strip() if primeira is not None else ""
    dados = DadosCurso(nome_pos)
    for nome_aba in abas:
        ws = _achar_aba(wb, nome_aba)
        if ws is None:
            raise ValueError(f"Aba '{nome_aba}' não encontrada na planilha.")
        aba = Aba(ws)
        col_nome, col_curso, col_email = detectar_colunas(aba)
        registros = []
        for linha in range(5, aba.ultima_linha(1) + 1):
            reg = ler_registro(aba, linha, col_nome, col_curso, col_email)
            if reg.rodape:
                continue
            registros.extend(dividir_por_docente(reg))
        letras = tuple(get_column_letter(c) for c in (col_nome, col_curso, col_email))
        dados.abas.append(DadosAba(ws.title, registros, letras))
    return dados


def aba_na_lista(lista, nome):
    """AbaNaLista."""
    return any(x.strip().casefold() == str(nome).casefold() for x in lista)
