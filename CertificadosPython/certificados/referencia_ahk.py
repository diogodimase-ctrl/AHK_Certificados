"""
Referência para conferência: monta o modelo-base do jeito do AHK
(CriarModeloBaseInterno da v27, chamada por chamada no PowerPoint) e compara
com o que o Python monta. Precisa do PowerPoint instalado.
"""

import os
import tempfile

from pptx import Presentation
from pptx.util import Emu

from . import modelo

TOLERANCIA_PT = 2.0


def _substituir_como_ahk(shape, marcador, valor, tamanho=0, negrito=-1):
    tr = shape.TextFrame.TextRange
    while True:
        pos = tr.Text.find(marcador) + 1
        if not pos:
            break
        chars = tr.Characters(pos, len(marcador))
        primeiro = chars.Characters(1, 1).Font
        nome, tam, neg, cor = primeiro.Name, primeiro.Size, primeiro.Bold, primeiro.Color.RGB
        chars.Text = valor
        novo = tr.Characters(pos, len(valor)).Font
        novo.Name, novo.Color.RGB = nome, cor
        novo.Size = tamanho if tamanho > 0 else tam
        novo.Bold = negrito if negrito >= 0 else neg


def _texto_com(shape):
    try:
        return shape.TextFrame.TextRange.Text
    except Exception:
        return None


def modelo_base_como_ahk(app, caminho_modelo, coordenador, cargo, data, assinatura, logo, destino):
    pres = app.Presentations.Open(os.path.abspath(caminho_modelo), False, False, False)
    try:
        slide = pres.Slides(1)
        for i in range(1, slide.Shapes.Count + 1):
            shape = slide.Shapes.Item(i)
            txt = _texto_com(shape)
            if txt is None:
                continue
            if "{COORDENADOR}" in txt or "{CARGO}" in txt or "{DATAEMISSAO}" in txt:
                _substituir_como_ahk(shape, "{COORDENADOR}", coordenador, 10, 0)
                _substituir_como_ahk(shape, "{CARGO}", cargo, 10, 1)
                _substituir_como_ahk(shape, "{DATAEMISSAO}", data)
            if "{COORDENADOR}" in txt or "{CARGO}" in txt:
                try:
                    shape.TextFrame.WordWrap = 0
                    shape.TextFrame2.AutoSize = 1
                except Exception:
                    pass
        formas = [slide.Shapes.Item(i) for i in range(1, slide.Shapes.Count + 1)]
        imagens = [s for s in formas if s.Type in (13, 7, 3)]
        if not imagens:
            imagens = [s for s in formas if not (_texto_com(s) or "").strip()]
        img_logo = img_assina = None
        if len(imagens) == 1:
            img_assina = imagens[0]
        elif len(imagens) >= 2:
            img_logo = min(imagens, key=lambda s: s.Top)
            img_assina = max(imagens, key=lambda s: s.Top)
        logo_top = 0
        if img_logo is not None:
            logo_top = img_logo.Top
            img_logo.Delete()
        if img_assina is not None:
            left, top = img_assina.Left, img_assina.Top
            img_assina.Delete()
            for i in range(1, slide.Shapes.Count + 1):
                sh = slide.Shapes.Item(i)
                if coordenador in (_texto_com(sh) or ""):
                    left, top = sh.Left, sh.Top - 50 - 8
                    break
            slide.Shapes.AddPicture(os.path.abspath(assinatura), False, True, left, top, 130, 50).Name = "ASSINATURA_REAL"
        if logo and os.path.isfile(logo):
            w = pres.PageSetup.SlideWidth
            slide.Shapes.AddPicture(os.path.abspath(logo), False, True, (w - 320) / 2,
                                    logo_top if logo_top > 0 else 20, 320, 200).Name = "LOGO_EINSTEIN"
        pres.SaveAs(os.path.abspath(destino), 24)
    finally:
        pres.Close()


def _formas(caminho):
    saida = {}
    for i, f in enumerate(Presentation(caminho).slides[0].shapes):
        if f.name in ("ASSINATURA_REAL", "LOGO_EINSTEIN"):
            chave = f.name
        elif f.has_text_frame and f.text_frame.text.strip():
            chave = " ".join(f.text_frame.text.split())[:40]
        else:
            chave = f"{f.name} #{i}"
        if f.left is None:
            continue
        saida[chave] = tuple(round(Emu(v).pt, 1) for v in (f.left, f.top, f.width, f.height))
    return saida


def comparar_modelo(caminho_modelo, coordenador, cargo, data, assinatura, logo):
    """Devolve (linhas, maior_diferença, (arquivo_ahk, arquivo_python), diagnóstico).
    Cada linha: (forma, ahk, python, ok)."""
    import comtypes
    import comtypes.client
    pasta = tempfile.mkdtemp(prefix="cert_comparar_")
    ahk, py = os.path.join(pasta, "ahk.pptx"), os.path.join(pasta, "python.pptx")
    comtypes.CoInitialize()
    try:
        app = comtypes.client.CreateObject("PowerPoint.Application")
        modelo_base_como_ahk(app, caminho_modelo, coordenador, cargo, data, assinatura, logo, ahk)
    finally:
        comtypes.CoUninitialize()
    diagnostico = []
    with open(py, "wb") as f:
        f.write(modelo.criar_modelo_base(caminho_modelo, coordenador, cargo, data, assinatura, logo,
                                         diagnostico=diagnostico))
    a, p = _formas(ahk), _formas(py)
    linhas, pior = [], 0.0
    for chave in sorted(set(a) | set(p)):
        va, vp = a.get(chave), p.get(chave)
        # Só esquerda, topo e largura: a altura das caixas que se ajustam ao
        # texto o PowerPoint recalcula ao abrir, sem mexer na posição do texto.
        dif = max(abs(x - y) for x, y in zip(va[:3], vp[:3])) if va and vp else float("inf")
        pior = max(pior, dif)
        linhas.append((chave, va, vp, dif <= TOLERANCIA_PT))
    return linhas, pior, (ahk, py), diagnostico
