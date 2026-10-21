"""
Preenchimento do modelo do certificado (.pptx) com python-pptx, sem abrir o
PowerPoint. Mesma lógica de CriarModeloBase e GerarCertificado do AHK:

1. Modelo-base do curso (uma vez por curso): troca {COORDENADOR}, {CARGO} e
   {DATAEMISSAO}, troca o logo e a assinatura.
2. Certificado (um por docente/aula): troca {NOME}, {POSGRADUACAO}, {CURSO},
   {DATA} e {HORAS} no modelo-base.
"""

import copy
import io
import os
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

# Tamanhos usados pelo AHK (em pontos).
ASSINATURA_LARGURA, ASSINATURA_ALTURA, ASSINATURA_MARGEM = 130, 50, 8
LOGO_LARGURA, LOGO_ALTURA, LOGO_TOPO_PADRAO = 320, 200, 20

_A_T = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
_TIPOS_IMAGEM = {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.CHART}


def _texto(forma):
    if not getattr(forma, "has_text_frame", False) or not forma.has_text_frame:
        return None
    return forma.text_frame.text


def _isolar_trecho(paragrafo, inicio, fim):
    """Deixa o trecho [inicio, fim) do parágrafo num pedaço (run) só e devolve
    esse run. O run novo herda a formatação do run onde o trecho começa (o AHK
    copia a formatação do primeiro caractere do marcador)."""
    runs = list(paragrafo.runs)
    textos = [r.text for r in runs]
    pos = 0
    i0 = i1 = o0 = o1 = None
    for i, t in enumerate(textos):
        if i0 is None and pos <= inicio < pos + len(t):
            i0, o0 = i, inicio - pos
        if pos < fim <= pos + len(t):
            i1, o1 = i, fim - pos
            break
        pos += len(t)

    antes = textos[i0][:o0]
    if i0 == i1:
        trecho, depois = textos[i0][o0:o1], textos[i0][o1:]
    else:
        trecho = textos[i0][o0:] + "".join(textos[i0 + 1:i1]) + textos[i1][:o1]
        depois = ""
        for k in range(i0 + 1, i1):
            runs[k].text = ""
        runs[i1].text = textos[i1][o1:]

    elemento = runs[i0]._r
    if depois:
        novo = copy.deepcopy(elemento)
        elemento.addnext(novo)
        novo.find(_A_T).text = depois
    meio = copy.deepcopy(elemento)
    elemento.addnext(meio)
    meio.find(_A_T).text = trecho
    if antes:
        elemento.find(_A_T).text = antes
    else:
        elemento.getparent().remove(elemento)
    return next(r for r in paragrafo.runs if r._r is meio)


def substituir(forma, marcador, valor, tamanho=None, negrito=None):
    """SubstituirPlaceholder: troca TODAS as ocorrências do marcador na forma,
    mantendo a formatação. tamanho/negrito forçam o tamanho e o negrito do
    texto novo (como no {COORDENADOR} e no {CARGO})."""
    if _texto(forma) is None:
        return 0
    trocas = 0
    for p in forma.text_frame.paragraphs:
        while True:
            texto = "".join(r.text for r in p.runs)
            pos = texto.find(marcador)
            if pos < 0:
                break
            run = _isolar_trecho(p, pos, pos + len(marcador))
            run.text = valor
            if tamanho:
                run.font.size = Pt(tamanho)
            if negrito is not None:
                run.font.bold = bool(negrito)
            trocas += 1
    return trocas


# ---------------------------------------------------------------------------
# "Ajustar forma ao texto" sem o PowerPoint
# ---------------------------------------------------------------------------
# No AHK, a caixa do coordenador recebe WordWrap = 0 e AutoSize = 1 e o
# PowerPoint encolhe a caixa na hora até a largura do texto (+ margens). Isso
# muda o Left da caixa, que é onde o AHK põe a assinatura. Medido no
# PowerPoint: alinhado à esquerda mantém a borda esquerda, centralizado mantém
# o centro, à direita mantém a borda direita.

_ARQUIVOS_DE_FONTE = None


def _arquivos_de_fonte():
    """{nome da fonte em minúsculas: arquivo}, das fontes do Windows e do usuário."""
    global _ARQUIVOS_DE_FONTE
    if _ARQUIVOS_DE_FONTE is None:
        _ARQUIVOS_DE_FONTE = {}
        try:
            import winreg
        except ImportError:
            return _ARQUIVOS_DE_FONTE
        pastas = [os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"),
                  os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "Windows", "Fonts")]
        for raiz in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            try:
                chave = winreg.OpenKey(raiz, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts")
            except OSError:
                continue
            i = 0
            while True:
                try:
                    nome, arquivo, _ = winreg.EnumValue(chave, i)
                except OSError:
                    break
                i += 1
                nome = re.sub(r"\s*\((TrueType|OpenType)\)$", "", nome).strip().lower()
                if not os.path.isabs(arquivo):
                    arquivo = next((os.path.join(p, arquivo) for p in pastas if os.path.isfile(os.path.join(p, arquivo))),
                                   os.path.join(pastas[0], arquivo))
                # "Arial & Arial Nova" etc.: vale para cada nome.
                for parte in nome.split(" & "):
                    _ARQUIVOS_DE_FONTE.setdefault(parte.strip(), arquivo)
    return _ARQUIVOS_DE_FONTE


def _arquivo_da_fonte(fonte, negrito):
    arquivos = _arquivos_de_fonte()
    nome = (fonte or "Calibri").lower()
    if negrito:
        arquivo = arquivos.get(f"{nome} bold") or arquivos.get(f"{nome} negrito")
        if arquivo:
            return arquivo, f"{fonte} Bold"
    if nome in arquivos:
        return arquivos[nome], fonte
    reserva = "calibri bold" if negrito else "calibri"
    return arquivos.get(reserva), f"Calibri (a fonte {fonte!r} não foi encontrada neste computador)"


def _medir(texto, fonte, tamanho, negrito):
    """Largura do texto em pontos (Pillow + arquivo da fonte do Windows)."""
    from PIL import ImageFont
    arquivo, _ = _arquivo_da_fonte(fonte, negrito)
    if arquivo and os.path.isfile(arquivo):
        try:
            return ImageFont.truetype(arquivo, tamanho * 20).getlength(texto) / 20
        except OSError:
            pass
    return len(texto) * tamanho * 0.5


_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
       "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


class EstilosDoTexto:
    """Descobre a fonte, o tamanho, o negrito e o alinhamento que valem para um
    texto, procurando nos mesmos lugares que o PowerPoint: o próprio texto, a
    caixa, o layout e o mestre (se for um espaço reservado), o estilo padrão da
    apresentação e, por fim, o tema."""

    def __init__(self, pres, slide):
        from lxml import etree
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        self.layout = slide.slide_layout
        self.mestre = self.layout.slide_master
        try:
            tema = etree.fromstring(self.mestre.part.part_related_by(RT.THEME).blob)
            self.tema = {"+mn-lt": tema.find(".//a:minorFont/a:latin", _NS).get("typeface"),
                         "+mj-lt": tema.find(".//a:majorFont/a:latin", _NS).get("typeface")}
        except Exception:
            self.tema = {"+mn-lt": "Calibri", "+mj-lt": "Calibri Light"}
        self.padrao = pres.part._element.find("p:defaultTextStyle", _NS)

    def _niveis(self, forma, nivel):
        """Elementos <a:lvlNpPr> do mais específico ao mais geral."""
        tag = f"a:lvl{nivel + 1}pPr"
        listas = [forma._element.find(".//a:lstStyle", _NS)]
        if forma.is_placeholder:
            ph = forma.placeholder_format
            for dono in (self.layout, self.mestre):
                alvo = next((s for s in dono.placeholders if s.placeholder_format.idx == ph.idx
                             or s.placeholder_format.type == ph.type), None)
                if alvo is not None:
                    listas.append(alvo._element.find(".//a:lstStyle", _NS))
            estilos = self.mestre._element.find("p:txStyles", _NS)
            if estilos is not None:
                nome = "p:titleStyle" if "TITLE" in str(ph.type) else "p:bodyStyle"
                listas.append(estilos.find(nome, _NS))
        else:
            listas.append(self.padrao)
        return [l.find(tag, _NS) for l in listas if l is not None and l.find(tag, _NS) is not None]

    def _procurar(self, rpr, niveis, atributo=None, filho=None):
        for el in [rpr] + [n.find("a:defRPr", _NS) for n in niveis]:
            if el is None:
                continue
            if filho is not None:
                f = el.find(filho, _NS)
                if f is not None and f.get("typeface"):
                    return f.get("typeface")
            elif el.get(atributo) is not None:
                return el.get(atributo)
        return None

    def do_texto(self, forma, paragrafo, run):
        niveis = self._niveis(forma, paragrafo.level)
        rpr = run._r.find("a:rPr", _NS)
        fonte = self._procurar(rpr, niveis, filho="a:latin") or "+mn-lt"
        fonte = self.tema.get(fonte, fonte)
        sz = self._procurar(rpr, niveis, "sz")
        b = self._procurar(rpr, niveis, "b")
        return fonte, (int(sz) / 100 if sz else 18.0), b in ("1", "true")

    def alinhamento(self, forma, paragrafo):
        if paragrafo.alignment is not None:
            return paragrafo.alignment
        for n in self._niveis(forma, paragrafo.level):
            if n.get("algn"):
                return {"ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "just": PP_ALIGN.JUSTIFY}.get(n.get("algn"), PP_ALIGN.LEFT)
        return PP_ALIGN.LEFT


def encolher_caixa(forma, estilos, diagnostico=None):
    """Faz o que o AutoSize do PowerPoint faria com a caixa (largura)."""
    tf = forma.text_frame
    larguras, alinhamento, usadas = [], None, []
    for p in tf.paragraphs:
        alinhamento = alinhamento or estilos.alinhamento(forma, p)
        total = 0.0
        for r in p.runs:
            fonte, tamanho, negrito = estilos.do_texto(forma, p, r)
            total += _medir(r.text, fonte, tamanho, negrito)
            descricao = f"{_arquivo_da_fonte(fonte, negrito)[1]} {tamanho:g} pt"
            if r.text.strip() and descricao not in usadas:
                usadas.append(descricao)
        larguras.append(total)
    margens = (tf.margin_left or 0) + (tf.margin_right or 0)
    nova = int(Pt(max(larguras or [0]))) + margens
    esquerda, largura = forma.left, forma.width
    if alinhamento == PP_ALIGN.CENTER:
        forma.left = int(esquerda + (largura - nova) / 2)
    elif alinhamento == PP_ALIGN.RIGHT:
        forma.left = int(esquerda + largura - nova)
    forma.width = nova
    if diagnostico is not None:
        nome_alin = {PP_ALIGN.CENTER: "centralizado", PP_ALIGN.RIGHT: "à direita"}.get(alinhamento, "à esquerda")
        diagnostico.append(f"Caixa do coordenador: texto {nome_alin}, medido com {', '.join(usadas)}; "
                           f"largura do texto {max(larguras or [0]):.1f} pt")


def _imagens_do_slide(slide):
    """Mesma escolha do AHK: imagens/objetos; se não houver, formas sem texto."""
    formas = list(slide.shapes)
    imagens = [f for f in formas if f.shape_type in _TIPOS_IMAGEM]
    if not imagens:
        imagens = [f for f in formas if not (_texto(f) or "").strip()]
    return imagens


def _remover(forma):
    forma._element.getparent().remove(forma._element)


def criar_modelo_base(caminho_modelo, coordenador, cargo, data_emissao, assinatura, logo, avisos=None,
                      diagnostico=None):
    """CriarModeloBase: devolve o modelo-base do curso, em bytes."""
    avisos = avisos if avisos is not None else []
    pres = Presentation(caminho_modelo)
    slide = pres.slides[0]
    estilos = EstilosDoTexto(pres, slide)

    for forma in slide.shapes:
        txt = _texto(forma)
        if txt is None:
            continue
        tinha_coord = "{COORDENADOR}" in txt or "{CARGO}" in txt
        if tinha_coord or "{DATAEMISSAO}" in txt:
            substituir(forma, "{COORDENADOR}", coordenador, 10, False)
            substituir(forma, "{CARGO}", cargo, 10, True)
            substituir(forma, "{DATAEMISSAO}", data_emissao)
        if tinha_coord:
            forma.text_frame.word_wrap = False
            forma.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
            encolher_caixa(forma, estilos, diagnostico)

    imagens = _imagens_do_slide(slide)
    logo_forma = assinatura_forma = None
    if len(imagens) == 1:
        assinatura_forma = imagens[0]
    elif len(imagens) >= 2:
        # Como no AHK: a mais alta é o logo, a mais baixa é a assinatura.
        logo_forma = min(imagens, key=lambda f: f.top)
        assinatura_forma = max(imagens, key=lambda f: f.top)
        if logo_forma is assinatura_forma:
            logo_forma = None

    logo_topo = 0
    if logo_forma is not None:
        logo_topo = logo_forma.top
        _remover(logo_forma)

    if assinatura_forma is not None:
        esquerda, topo = assinatura_forma.left, assinatura_forma.top
        _remover(assinatura_forma)
        coord_forma = next((f for f in slide.shapes if coordenador in (_texto(f) or "")), None)
        if coord_forma is not None:
            esquerda = coord_forma.left
            topo = coord_forma.top - Pt(ASSINATURA_ALTURA + ASSINATURA_MARGEM)
        slide.shapes.add_picture(assinatura, esquerda, topo, Pt(ASSINATURA_LARGURA), Pt(ASSINATURA_ALTURA)).name = "ASSINATURA_REAL"
    else:
        avisos.append("Nenhuma imagem de assinatura encontrada no modelo: o certificado sai sem assinatura.")

    if logo and os.path.isfile(logo):
        esquerda = int((pres.slide_width - Pt(LOGO_LARGURA)) / 2)
        topo = logo_topo if logo_topo > 0 else Pt(LOGO_TOPO_PADRAO)
        slide.shapes.add_picture(logo, esquerda, topo, Pt(LOGO_LARGURA), Pt(LOGO_ALTURA)).name = "LOGO_EINSTEIN"
    else:
        avisos.append(f"Logo não encontrado ({logo}): o certificado sai sem o logo.")

    saida = io.BytesIO()
    pres.save(saida)
    return saida.getvalue()


def preencher_certificado(modelo_base, nome, pos_graduacao, curso, data, horas, destino):
    """GerarCertificado (parte do PowerPoint): grava o .pptx do certificado."""
    pres = Presentation(io.BytesIO(modelo_base))
    for forma in pres.slides[0].shapes:
        if "{NOME}" not in (_texto(forma) or ""):
            continue
        substituir(forma, "{NOME}", nome)
        substituir(forma, "{POSGRADUACAO}", pos_graduacao)
        substituir(forma, "{CURSO}", curso)
        substituir(forma, "{DATA}", data)
        substituir(forma, "{HORAS}", horas)
    pres.save(destino)
