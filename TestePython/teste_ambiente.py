"""
Teste de ambiente — Certificados em Python
==========================================

Programa pequeno para rodar UMA vez no computador do trabalho e descobrir se
o caminho "B" (app Python no lugar do AHK) é viável lá. Ele NÃO envia e-mail,
NÃO mexe na pasta "Para Enviar" e NÃO altera nenhuma planilha: só lê os
arquivos que você escolher e grava resultados em uma pasta de teste.

O que ele verifica:
  1. Se o próprio .exe consegue rodar (se você está vendo a janela, passou).
  2. Se consegue gravar arquivos no computador e no OneDrive da empresa.
  3. Se o python-pptx preenche um modelo .pptx (um de exemplo e, se você
     escolher, o seu modelo real) sem precisar abrir o PowerPoint.
  4. Se consegue trocar a imagem da assinatura (opcional).
  5. Se o PowerPoint instalado aceita ser chamado só para exportar o PDF.
  6. Se o LibreOffice está instalado (outra forma de gerar PDF).
  7. Se o openpyxl lê uma planilha real (opcional).
  8. Gera uma versão HTML do certificado, usada no teste 2 (Power Automate).

Tudo vai para a pasta "Teste_Certificados_Python" (no OneDrive da empresa,
se encontrar; senão na Área de Trabalho), junto com um relatorio.txt.
"""

import os
import re
import sys
import time
import shutil
import struct
import zlib
import base64
import platform
import tempfile
import threading
import traceback
import subprocess
from datetime import datetime

import tkinter as tk
from tkinter import ttk, filedialog

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


VERSAO = "1.0"
NOME_PASTA = "Teste_Certificados_Python"

# Valores fictícios usados para preencher os modelos.
VALORES_TESTE = {
    "{NOME}": "Dra. Maria Exemplo da Silva",
    "{POSGRADUACAO}": "Pós-Graduação em Teste de Sistemas",
    "{CURSO}": "Aula de Demonstração",
    "{DATA}": "20/10/2026",
    "{HORAS}": "2 horas",
    "{COORDENADOR}": "Dr. Fulano Coordenador",
    "{CARGO}": "Coordenador",
    "{DATAEMISSAO}": "São Paulo, 20 de outubro de 2026",
}

OK, FALHOU, PULADO, AVISO = "OK", "FALHOU", "PULADO", "AVISO"


# ---------------------------------------------------------------------------
# Pasta de saída
# ---------------------------------------------------------------------------

def pasta_onedrive_empresa():
    """Pasta do OneDrive corporativo (a que sincroniza com o SharePoint)."""
    caminho = os.environ.get("OneDriveCommercial")
    if caminho and os.path.isdir(caminho):
        return caminho
    # Às vezes a variável não existe, mas a pasta "OneDrive - <Empresa>" sim.
    perfil = os.environ.get("USERPROFILE", "")
    try:
        for nome in os.listdir(perfil):
            if nome.startswith("OneDrive - ") and os.path.isdir(os.path.join(perfil, nome)):
                return os.path.join(perfil, nome)
    except OSError:
        pass
    return None


def pasta_area_de_trabalho():
    perfil = os.environ.get("USERPROFILE", "")
    for nome in ("Desktop", "Área de Trabalho"):
        caminho = os.path.join(perfil, nome)
        if os.path.isdir(caminho):
            return caminho
    return perfil


# ---------------------------------------------------------------------------
# Troca de {MARCADORES} no .pptx
# ---------------------------------------------------------------------------

def substituir_no_paragrafo(paragrafo, mapa):
    """Troca os marcadores mesmo quando o PowerPoint quebrou o texto em vários
    pedaços ("runs"). O texto novo fica com a formatação do primeiro pedaço,
    como o AHK já faz. Devolve a lista de marcadores trocados."""
    runs = list(paragrafo.runs)
    if not runs:
        return []
    trocados = []
    textos = [r.text for r in runs]
    for marcador, valor in mapa.items():
        while marcador in "".join(textos):
            inicio = "".join(textos).index(marcador)
            fim = inicio + len(marcador)
            # Descobre em qual pedaço o marcador começa e em qual termina.
            pos = 0
            i0 = i1 = o0 = o1 = None
            for i, t in enumerate(textos):
                if i0 is None and pos <= inicio < pos + len(t):
                    i0, o0 = i, inicio - pos
                if pos < fim <= pos + len(t):
                    i1, o1 = i, fim - pos
                    break
                pos += len(t)
            if i0 == i1:
                textos[i0] = textos[i0][:o0] + valor + textos[i0][o1:]
            else:
                textos[i0] = textos[i0][:o0] + valor
                for k in range(i0 + 1, i1):
                    textos[k] = ""
                textos[i1] = textos[i1][o1:]
            trocados.append(marcador)
    for r, t in zip(runs, textos):
        if r.text != t:
            r.text = t
    return trocados


def caixas_de_texto(formas):
    """Todas as caixas de texto do slide, inclusive dentro de grupos e tabelas."""
    for forma in formas:
        if forma.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from caixas_de_texto(forma.shapes)
        elif getattr(forma, "has_table", False) and forma.has_table:
            for linha in forma.table.rows:
                for celula in linha.cells:
                    yield celula.text_frame
        elif getattr(forma, "has_text_frame", False) and forma.has_text_frame:
            yield forma.text_frame


def marcadores_no_modelo(pres):
    """Lista os {MARCADORES} do modelo e quais estão quebrados em pedaços."""
    achados, quebrados = set(), set()
    for slide in pres.slides:
        for tf in caixas_de_texto(slide.shapes):
            for p in tf.paragraphs:
                texto = "".join(r.text for r in p.runs)
                for m in re.findall(r"\{[A-ZÇÃÕÁÉÍÓÚ_]+\}", texto):
                    achados.add(m)
                    if not any(m in r.text for r in p.runs):
                        quebrados.add(m)
    return sorted(achados), sorted(quebrados)


def preencher_pptx(origem, destino, mapa):
    pres = Presentation(origem)
    trocados = []
    for slide in pres.slides:
        for tf in caixas_de_texto(slide.shapes):
            for p in tf.paragraphs:
                trocados += substituir_no_paragrafo(p, mapa)
    pres.save(destino)
    return trocados


def trocar_assinatura(caminho_pptx, caminho_imagem):
    """Mesma regra do AHK: com 2+ imagens, a mais baixa é a assinatura.
    Apaga essa imagem e põe a nova no mesmo lugar (130 x 50 pontos)."""
    pres = Presentation(caminho_pptx)
    slide = pres.slides[0]
    imagens = [f for f in slide.shapes if f.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if not imagens:
        return "o modelo não tem nenhuma imagem para trocar"
    assinatura = max(imagens, key=lambda f: f.top)
    esquerda, topo = assinatura.left, assinatura.top
    assinatura._element.getparent().remove(assinatura._element)
    slide.shapes.add_picture(caminho_imagem, esquerda, topo, width=Pt(130), height=Pt(50))
    pres.save(caminho_pptx)
    return None


# ---------------------------------------------------------------------------
# Modelo de exemplo (criado na hora, para não depender de arquivo nenhum)
# ---------------------------------------------------------------------------

def png_liso(largura, altura, rgb):
    """PNG de uma cor só, montado à mão (sem biblioteca de imagem)."""
    linha = b"\x00" + bytes(rgb) * largura
    dados = zlib.compress(linha * altura)

    def bloco(tipo, conteudo):
        return (struct.pack(">I", len(conteudo)) + tipo + conteudo
                + struct.pack(">I", zlib.crc32(tipo + conteudo) & 0xFFFFFFFF))

    return (b"\x89PNG\r\n\x1a\n"
            + bloco(b"IHDR", struct.pack(">IIBBBBB", largura, altura, 8, 2, 0, 0, 0))
            + bloco(b"IDAT", dados) + bloco(b"IEND", b""))


def criar_modelo_exemplo(destino, pasta_temp):
    pres = Presentation()
    pres.slide_width, pres.slide_height = Cm(29.7), Cm(21)   # A4 deitado
    slide = pres.slides.add_slide(pres.slide_layouts[6])     # slide em branco

    faixa = os.path.join(pasta_temp, "faixa.png")
    with open(faixa, "wb") as f:
        f.write(png_liso(600, 40, (15, 107, 111)))
    slide.shapes.add_picture(faixa, Cm(2), Cm(1.5), width=Cm(25.7), height=Cm(1.6))

    def caixa(texto, topo, tamanho, negrito=False, cor=(20, 38, 46)):
        tb = slide.shapes.add_textbox(Cm(2), topo, Cm(25.7), Cm(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        # Divide o texto em pedaços de propósito, para testar marcador quebrado.
        for pedaco in re.split(r"(\{[A-Z]+)", texto):
            if not pedaco:
                continue
            r = p.add_run()
            r.text = pedaco
            r.font.size = Pt(tamanho)
            r.font.bold = negrito
            r.font.color.rgb = RGBColor(*cor)

    caixa("DECLARAÇÃO", Cm(4), 32, True, (15, 107, 111))
    caixa("Declaramos que {NOME} ministrou a aula {CURSO} na {POSGRADUACAO}, "
          "em {DATA}, com carga horária de {HORAS}.", Cm(7), 16)
    caixa("{DATAEMISSAO}", Cm(12), 12)

    assinatura = os.path.join(pasta_temp, "assinatura_exemplo.png")
    with open(assinatura, "wb") as f:
        f.write(png_liso(260, 100, (200, 200, 200)))
    slide.shapes.add_picture(assinatura, Cm(12.56), Cm(14.2), width=Pt(130), height=Pt(50))
    caixa("{COORDENADOR}", Cm(16.3), 11, True)
    caixa("{CARGO}", Cm(17.2), 10)
    pres.save(destino)


def criar_html(destino):
    """Versão HTML do certificado (teste 2: converter HTML em PDF na nuvem).
    Layout com tabela e estilos simples, que conversores antigos entendem."""
    faixa = base64.b64encode(png_liso(600, 40, (15, 107, 111))).decode()
    v = VALORES_TESTE
    html = f"""<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="utf-8"><title>Certificado de teste</title>
<style>
@page {{ size: A4 landscape; margin: 1.5cm; }}
body {{ font-family: Calibri, Arial, sans-serif; color: #14262e; }}
h1 {{ color: #0f6b6f; font-size: 32pt; text-align: center; margin: 24pt 0; }}
p {{ font-size: 16pt; text-align: center; line-height: 1.5; }}
.pequeno {{ font-size: 11pt; }}
</style></head>
<body>
<img src="data:image/png;base64,{faixa}" width="100%" height="40" alt="faixa">
<h1>DECLARAÇÃO</h1>
<p>Declaramos que <b>{v['{NOME}']}</b> ministrou a aula <b>{v['{CURSO}']}</b> na
{v['{POSGRADUACAO}']}, em {v['{DATA}']}, com carga horária de {v['{HORAS}']}.</p>
<p class="pequeno">{v['{DATAEMISSAO}']}</p>
<table width="100%" style="margin-top:48pt"><tr><td align="center">
<div style="width:180pt;height:50pt;background:#c8c8c8;margin:0 auto"></div>
<p class="pequeno"><b>{v['{COORDENADOR}']}</b><br>{v['{CARGO}']}</p>
</td></tr></table>
<p class="pequeno">Acentos: ação, coração, pós-graduação, São Paulo — ç ã õ é ê í ó ú</p>
</body></html>"""
    with open(destino, "w", encoding="utf-8") as f:
        f.write(html)


# ---------------------------------------------------------------------------
# PDF pelo PowerPoint instalado
# ---------------------------------------------------------------------------

def powerpoint_ja_aberto():
    try:
        saida = subprocess.run(["tasklist", "/FI", "IMAGENAME eq POWERPNT.EXE"],
                               capture_output=True, text=True,
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        return "POWERPNT.EXE" in saida.stdout.upper()
    except Exception:
        return False


def exportar_pdf_powerpoint(lista_pptx):
    """Abre cada .pptx escondido e salva como PDF. Só fecha o PowerPoint se
    ele não estava aberto antes (para não fechar o trabalho de ninguém)."""
    import comtypes
    import comtypes.client
    comtypes.CoInitialize()
    estava_aberto = powerpoint_ja_aberto()
    app = comtypes.client.CreateObject("PowerPoint.Application")
    tempos = []
    try:
        for pptx in lista_pptx:
            pdf = os.path.splitext(pptx)[0] + "_PowerPoint.pdf"
            inicio = time.time()
            pres = app.Presentations.Open(os.path.abspath(pptx), True, False, False)
            pres.SaveAs(os.path.abspath(pdf), 32)   # 32 = PDF
            pres.Close()
            tempos.append((os.path.basename(pdf), time.time() - inicio))
    finally:
        if not estava_aberto:
            try:
                if app.Presentations.Count == 0:
                    app.Quit()
            except Exception:
                pass
        comtypes.CoUninitialize()
    return tempos


def achar_libreoffice():
    for base in (os.environ.get("ProgramFiles", ""), os.environ.get("ProgramFiles(x86)", "")):
        caminho = os.path.join(base, "LibreOffice", "program", "soffice.exe")
        if base and os.path.isfile(caminho):
            return caminho
    return shutil.which("soffice")


# ---------------------------------------------------------------------------
# Planilha
# ---------------------------------------------------------------------------

def ler_planilha(caminho, pasta_temp):
    import openpyxl
    copia = os.path.join(pasta_temp, "planilha_copia.xlsx")
    shutil.copy2(caminho, copia)   # lê uma cópia: o original fica intocado
    inicio = time.time()
    wb = openpyxl.load_workbook(copia, data_only=True)
    linhas = [f"{len(wb.sheetnames)} aba(s): " + ", ".join(wb.sheetnames[:8])
              + (" …" if len(wb.sheetnames) > 8 else "")]
    ws = wb[wb.sheetnames[0]]
    linhas.append(f"1ª aba '{ws.title}': {ws.max_row} linhas × {ws.max_column} colunas, "
                  f"{len(ws.merged_cells.ranges)} células mescladas")
    cab = [str(c.value).strip() for c in ws[3] if c.value not in (None, "")] if ws.max_row >= 3 else []
    if cab:
        linhas.append("Linha 3: " + " | ".join(cab[:8]))
    linhas.append(f"Tempo de leitura: {time.time() - inicio:.1f} s")
    wb.close()
    return linhas


# ---------------------------------------------------------------------------
# Janela
# ---------------------------------------------------------------------------

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"Teste de ambiente — Certificados em Python (v{VERSAO})")
        self.geometry("860x640")
        self.minsize(720, 520)
        self.modelo_real = None
        self.assinatura = None
        self.planilha = None
        self.resultados = []
        self.pasta_saida = None
        self._montar()

    def _montar(self):
        estilo = ttk.Style(self)
        try:
            estilo.theme_use("vista")
        except tk.TclError:
            pass
        estilo.configure("Titulo.TLabel", font=("Segoe UI", 15, "bold"))
        estilo.configure("Texto.TLabel", font=("Segoe UI", 10))
        estilo.configure("Grande.TButton", font=("Segoe UI", 11, "bold"), padding=8)

        topo = ttk.Frame(self, padding=(16, 14, 16, 6))
        topo.pack(fill="x")
        ttk.Label(topo, text="Teste de ambiente — Certificados em Python", style="Titulo.TLabel").pack(anchor="w")
        ttk.Label(topo, style="Texto.TLabel", wraplength=800, justify="left", text=(
            "Se você está vendo esta janela, o primeiro teste já passou: o computador deixou o programa rodar.\n"
            "Os itens abaixo são opcionais, mas ajudam muito. Nada é enviado e nenhum arquivo original é alterado.")
        ).pack(anchor="w", pady=(4, 0))

        opc = ttk.LabelFrame(self, text=" Opcionais (escolha antes de rodar) ", padding=10)
        opc.pack(fill="x", padx=16, pady=8)
        self.lbl_modelo = self._linha_opcional(opc, 0, "Meu modelo de certificado (.pptx)", self.escolher_modelo)
        self.lbl_assin = self._linha_opcional(opc, 1, "Uma imagem de assinatura (.png/.jpg)", self.escolher_assinatura)
        self.lbl_plan = self._linha_opcional(opc, 2, "Uma planilha de curso ou a mestre (.xlsx)", self.escolher_planilha)

        botoes = ttk.Frame(self, padding=(16, 0))
        botoes.pack(fill="x")
        self.btn_rodar = ttk.Button(botoes, text="Rodar testes", style="Grande.TButton", command=self.rodar)
        self.btn_rodar.pack(side="left")
        self.btn_pasta = ttk.Button(botoes, text="Abrir pasta dos resultados", command=self.abrir_pasta, state="disabled")
        self.btn_pasta.pack(side="left", padx=8)
        self.btn_copiar = ttk.Button(botoes, text="Copiar relatório", command=self.copiar, state="disabled")
        self.btn_copiar.pack(side="left")

        tabela = ttk.Frame(self, padding=(16, 10, 16, 16))
        tabela.pack(fill="both", expand=True)
        self.arvore = ttk.Treeview(tabela, columns=("status", "detalhe"), show="tree headings", height=12)
        self.arvore.heading("#0", text="Teste")
        self.arvore.heading("status", text="Resultado")
        self.arvore.heading("detalhe", text="Detalhe")
        self.arvore.column("#0", width=250, stretch=False)
        self.arvore.column("status", width=90, stretch=False, anchor="center")
        self.arvore.column("detalhe", width=460)
        self.arvore.tag_configure(OK, foreground="#2e7d4f")
        self.arvore.tag_configure(FALHOU, foreground="#b03a3a")
        self.arvore.tag_configure(AVISO, foreground="#9a6a00")
        self.arvore.tag_configure(PULADO, foreground="#5a6b72")
        rolagem = ttk.Scrollbar(tabela, orient="vertical", command=self.arvore.yview)
        self.arvore.configure(yscrollcommand=rolagem.set)
        self.arvore.pack(side="left", fill="both", expand=True)
        rolagem.pack(side="right", fill="y")

        self.status = ttk.Label(self, text="Pronto para rodar.", style="Texto.TLabel", padding=(16, 0, 16, 10))
        self.status.pack(fill="x")

    def _linha_opcional(self, pai, linha, texto, comando):
        ttk.Button(pai, text="Escolher…", command=comando).grid(row=linha, column=0, sticky="w", pady=2)
        ttk.Label(pai, text=texto, style="Texto.TLabel").grid(row=linha, column=1, sticky="w", padx=8)
        lbl = ttk.Label(pai, text="(não escolhido)", foreground="#5a6b72")
        lbl.grid(row=linha, column=2, sticky="w")
        return lbl

    def escolher_modelo(self):
        c = filedialog.askopenfilename(title="Modelo de certificado", filetypes=[("PowerPoint", "*.pptx")])
        if c:
            self.modelo_real = c
            self.lbl_modelo.config(text=os.path.basename(c))

    def escolher_assinatura(self):
        c = filedialog.askopenfilename(title="Imagem de assinatura",
                                       filetypes=[("Imagens", "*.png *.jpg *.jpeg *.gif *.bmp")])
        if c:
            self.assinatura = c
            self.lbl_assin.config(text=os.path.basename(c))

    def escolher_planilha(self):
        c = filedialog.askopenfilename(title="Planilha", filetypes=[("Excel", "*.xlsx *.xlsm")])
        if c:
            self.planilha = c
            self.lbl_plan.config(text=os.path.basename(c))

    # -- resultados ---------------------------------------------------------

    def registrar(self, teste, status, detalhe=""):
        self.resultados.append((teste, status, detalhe))
        self.after(0, lambda: self.arvore.insert("", "end", text=teste, values=(status, detalhe), tags=(status,)))

    def avisar(self, texto):
        self.after(0, lambda: self.status.config(text=texto))

    def rodar(self):
        self.btn_rodar.config(state="disabled")
        self.arvore.delete(*self.arvore.get_children())
        self.resultados = []
        threading.Thread(target=self._rodar_testes, daemon=True).start()

    def _rodar_testes(self):
        temp = tempfile.mkdtemp(prefix="teste_cert_")
        gerados_pptx = []
        try:
            self.registrar("Programa abriu", OK, f"Windows: {platform.platform()}")

            # 1. Pasta de saída
            od = pasta_onedrive_empresa()
            base = od or pasta_area_de_trabalho()
            self.pasta_saida = os.path.join(base, NOME_PASTA)
            try:
                os.makedirs(self.pasta_saida, exist_ok=True)
                with open(os.path.join(self.pasta_saida, "gravacao.txt"), "w", encoding="utf-8") as f:
                    f.write("teste de gravação\n")
                onde = "OneDrive da empresa" if od else "Área de Trabalho (OneDrive da empresa não encontrado)"
                self.registrar("Gravar arquivos", OK if od else AVISO, f"{onde}: {self.pasta_saida}")
            except Exception as e:
                self.pasta_saida = os.path.join(temp, NOME_PASTA)
                os.makedirs(self.pasta_saida, exist_ok=True)
                self.registrar("Gravar arquivos", FALHOU, f"{e} — usando pasta temporária")

            # 2. Modelo de exemplo
            self.avisar("Preenchendo o modelo de exemplo…")
            try:
                modelo_ex = os.path.join(self.pasta_saida, "modelo_exemplo.pptx")
                criar_modelo_exemplo(modelo_ex, temp)
                saida = os.path.join(self.pasta_saida, "certificado_teste.pptx")
                inicio = time.time()
                trocados = preencher_pptx(modelo_ex, saida, VALORES_TESTE)
                restantes, _ = marcadores_no_modelo(Presentation(saida))
                status = OK if not restantes else FALHOU
                self.registrar("python-pptx (exemplo)", status,
                               f"{len(trocados)} marcadores trocados em {time.time() - inicio:.2f} s"
                               + (f"; sobraram {restantes}" if restantes else ""))
                gerados_pptx.append(saida)
            except Exception as e:
                self.registrar("python-pptx (exemplo)", FALHOU, str(e))

            # 3. Modelo real
            if self.modelo_real:
                self.avisar("Preenchendo o seu modelo…")
                try:
                    achados, quebrados = marcadores_no_modelo(Presentation(self.modelo_real))
                    saida = os.path.join(self.pasta_saida, "certificado_modelo_real.pptx")
                    preencher_pptx(self.modelo_real, saida, VALORES_TESTE)
                    restantes, _ = marcadores_no_modelo(Presentation(saida))
                    desconhecidos = [m for m in achados if m not in VALORES_TESTE]
                    detalhe = f"marcadores no modelo: {', '.join(achados) or 'nenhum'}"
                    if quebrados:
                        detalhe += f"; quebrados em pedaços (tratados): {', '.join(quebrados)}"
                    if desconhecidos:
                        detalhe += f"; não conhecidos pelo teste: {', '.join(desconhecidos)}"
                    ok = achados and not [m for m in restantes if m in VALORES_TESTE]
                    self.registrar("python-pptx (seu modelo)", OK if ok else AVISO, detalhe)
                    gerados_pptx.append(saida)
                    if self.assinatura:
                        erro = trocar_assinatura(saida, self.assinatura)
                        self.registrar("Trocar assinatura", AVISO if erro else OK,
                                       erro or "assinatura trocada no certificado_modelo_real.pptx")
                except Exception as e:
                    self.registrar("python-pptx (seu modelo)", FALHOU, str(e))
            else:
                self.registrar("python-pptx (seu modelo)", PULADO, "nenhum modelo escolhido")
                if self.assinatura:
                    self.registrar("Trocar assinatura", PULADO, "precisa do seu modelo também")

            # 4. PDF pelo PowerPoint
            if gerados_pptx:
                self.avisar("Pedindo ao PowerPoint para salvar em PDF…")
                try:
                    tempos = exportar_pdf_powerpoint(gerados_pptx)
                    self.registrar("PDF pelo PowerPoint", OK,
                                   "; ".join(f"{n} em {t:.1f} s" for n, t in tempos))
                except Exception as e:
                    self.registrar("PDF pelo PowerPoint", FALHOU, f"{type(e).__name__}: {e}")
            else:
                self.registrar("PDF pelo PowerPoint", PULADO, "nenhum .pptx foi gerado")

            # 5. LibreOffice
            lo = achar_libreoffice()
            self.registrar("LibreOffice instalado", OK if lo else AVISO, lo or "não encontrado (não é obrigatório)")

            # 6. HTML para o teste 2
            try:
                criar_html(os.path.join(self.pasta_saida, "certificado_teste.html"))
                self.registrar("HTML para o teste 2", OK, "certificado_teste.html criado")
            except Exception as e:
                self.registrar("HTML para o teste 2", FALHOU, str(e))

            # 7. Planilha
            if self.planilha:
                self.avisar("Lendo a planilha…")
                try:
                    linhas = ler_planilha(self.planilha, temp)
                    self.registrar("Ler planilha (openpyxl)", OK, linhas[0])
                    for extra in linhas[1:]:
                        self.registrar("   ", OK, extra)
                except Exception as e:
                    self.registrar("Ler planilha (openpyxl)", FALHOU, f"{type(e).__name__}: {e}")
            else:
                self.registrar("Ler planilha (openpyxl)", PULADO, "nenhuma planilha escolhida")

        except Exception:
            self.registrar("Erro inesperado", FALHOU, traceback.format_exc().splitlines()[-1])
        finally:
            shutil.rmtree(temp, ignore_errors=True)
            self._salvar_relatorio()
            self.after(0, self._fim)

    def texto_relatorio(self):
        linhas = [f"Teste de ambiente — Certificados em Python v{VERSAO}",
                  f"Data: {datetime.now():%d/%m/%Y %H:%M}", f"Pasta: {self.pasta_saida}", ""]
        for teste, status, detalhe in self.resultados:
            linhas.append(f"[{status:^7}] {teste.strip() or '…'} — {detalhe}")
        return "\n".join(linhas)

    def _salvar_relatorio(self):
        try:
            with open(os.path.join(self.pasta_saida, "relatorio.txt"), "w", encoding="utf-8") as f:
                f.write(self.texto_relatorio() + "\n")
        except Exception:
            pass

    def _fim(self):
        falhas = sum(1 for _, s, _ in self.resultados if s == FALHOU)
        self.status.config(text=("Terminou sem falhas. " if not falhas else f"Terminou com {falhas} falha(s). ")
                           + "O relatório está em relatorio.txt, na pasta dos resultados.")
        self.btn_rodar.config(state="normal")
        self.btn_pasta.config(state="normal")
        self.btn_copiar.config(state="normal")

    def abrir_pasta(self):
        if self.pasta_saida and os.path.isdir(self.pasta_saida):
            os.startfile(self.pasta_saida)

    def copiar(self):
        self.clipboard_clear()
        self.clipboard_append(self.texto_relatorio())
        self.status.config(text="Relatório copiado. Pode colar aqui no chat.")


if __name__ == "__main__":
    app = App()
    # "--automatico": roda sozinho e fecha (usado para conferir o .exe antes de entregar).
    if "--automatico" in sys.argv:
        app._fim_original = app._fim
        app._fim = lambda: (app._fim_original(), app.after(500, app.destroy))
        app.after(500, app.rodar)
    app.mainloop()
